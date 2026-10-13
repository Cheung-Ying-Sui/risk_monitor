"""Build a clean, data-free template from the supplied questionnaire's labels.

Usage: python scripts/build_sanctions_report_template.py /path/to/reference.docx
The source is read only. No sample answers, legacy logos or custom XML are copied.
"""
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def build(source: str) -> None:
    reference = Document(source)
    labels = [row.cells[0].text for row in reference.tables[1].rows]
    if len(labels) != 24:
        raise ValueError("Expected the 24-row bilingual questionnaire")
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.54)
    section.left_margin = section.right_margin = Cm(2.49)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(9)
    normal.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    normal.paragraph_format.space_after = Pt(4)
    for name in ('Title', 'Heading 1', 'Heading 2'):
        document.styles[name].element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体')
    document.add_paragraph('Information 风险信息')
    document.add_paragraph('制裁风险尽职调查问卷', style='Title')
    document.add_paragraph('SANCTIONS DUE DILIGENCE QUESTIONNAIRE')
    intro = document.add_table(rows=1, cols=1)
    intro.style = 'Table Grid'
    intro.cell(0, 0).text = ('The following information supports the assessment of the risk at the time of underwriting.\n'
                             '以下申报信息供核保风险评估参考；未采集的信息标记为“未提供”。')
    document.add_paragraph()
    table = document.add_table(rows=len(labels), cols=2)
    table.style, table.autofit = 'Table Grid', False
    for row, label in zip(table.rows, labels):
        row.cells[0].text = label
        row.cells[0].width, row.cells[1].width = Cm(4.5), Cm(11.52)
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    table.columns[0].width, table.columns[1].width = Cm(4.5), Cm(11.52)
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    document.core_properties.title = '制裁风险尽职调查问卷'
    for rel_id, rel in list(document.part.rels.items()):
        if rel.reltype.endswith('/customXml'):
            document.part.drop_rel(rel_id)
    destination = Path(__file__).resolve().parents[1] / 'risk_monitor/assets/sanctions/questionnaire-template.docx'
    document.save(destination)


if __name__ == '__main__':
    build(sys.argv[1])
