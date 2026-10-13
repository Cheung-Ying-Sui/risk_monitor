from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable


PARSER_NAME = "marine_cargo_excel"
PARSER_VERSION = "1.3.0"
SUPPORTED_EXTENSIONS = {".xls", ".xlsx"}


class InsuranceWorkbookParseError(ValueError):
    pass


@dataclass(frozen=True)
class GridCell:
    row: int
    column: int
    value: Any
    number_format: str | None = None


@dataclass(frozen=True)
class SheetGrid:
    name: str
    cells: tuple[GridCell, ...]
    max_row: int
    max_column: int

    def value(self, row: int, column: int) -> Any:
        for cell in self.cells:
            if cell.row == row and cell.column == column:
                return cell.value
        return None

    def cell(self, row: int, column: int) -> GridCell | None:
        for cell in self.cells:
            if cell.row == row and cell.column == column:
                return cell
        return None

    def non_empty(self) -> Iterable[GridCell]:
        return (cell for cell in self.cells if _clean_text(cell.value))


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def _normalized_label(value: Any) -> str:
    text = _clean_text(value) or ""
    return re.sub(r"[\s:：._/()-]+", "", text).lower()


def _decimal(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = _clean_text(value)
    if not text:
        return None
    match = re.search(r"[-+]?\d[\d,]*(?:\.\d+)?", text)
    if not match:
        return None
    try:
        return Decimal(match.group(0).replace(",", ""))
    except InvalidOperation:
        return None


def _datetime(value: Any) -> datetime | None:
    if isinstance(value, datetime):
        return value
    text = _clean_text(value)
    if not text:
        return None
    for pattern in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%Y.%m.%d"):
        try:
            return datetime.strptime(text, pattern)
        except ValueError:
            continue
    return None


def _load_xlsx(content: bytes) -> list[SheetGrid]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise InsuranceWorkbookParseError("openpyxl is required to parse .xlsx files") from exc

    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True, read_only=False)
    except Exception as exc:
        raise InsuranceWorkbookParseError("Unable to read the .xlsx workbook") from exc

    sheets: list[SheetGrid] = []
    for worksheet in workbook.worksheets:
        cells = tuple(
            GridCell(
                row=cell.row,
                column=cell.column,
                value=cell.value,
                number_format=cell.number_format,
            )
            for row in worksheet.iter_rows()
            for cell in row
            if cell.value is not None
        )
        sheets.append(
            SheetGrid(
                name=worksheet.title,
                cells=cells,
                max_row=worksheet.max_row,
                max_column=worksheet.max_column,
            )
        )
    return sheets


def _load_xls(content: bytes) -> list[SheetGrid]:
    try:
        import xlrd
    except ImportError as exc:
        raise InsuranceWorkbookParseError("xlrd is required to parse .xls files") from exc

    try:
        workbook = xlrd.open_workbook(file_contents=content, formatting_info=True)
    except Exception as exc:
        raise InsuranceWorkbookParseError("Unable to read the .xls workbook") from exc

    sheets: list[SheetGrid] = []
    for worksheet in workbook.sheets():
        cells: list[GridCell] = []
        for row_index in range(worksheet.nrows):
            for column_index in range(worksheet.ncols):
                source = worksheet.cell(row_index, column_index)
                if source.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                    continue
                value: Any = source.value
                if source.ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate.xldate_as_datetime(source.value, workbook.datemode)
                number_format = None
                try:
                    xf = workbook.xf_list[source.xf_index]
                    number_format = workbook.format_map[xf.format_key].format_str
                except (AttributeError, IndexError, KeyError):
                    pass
                cells.append(
                    GridCell(
                        row=row_index + 1,
                        column=column_index + 1,
                        value=value,
                        number_format=number_format,
                    )
                )
        sheets.append(
            SheetGrid(
                name=worksheet.name,
                cells=tuple(cells),
                max_row=worksheet.nrows,
                max_column=worksheet.ncols,
            )
        )
    return sheets


def _row_cells(sheet: SheetGrid, row: int) -> list[GridCell]:
    return sorted((cell for cell in sheet.cells if cell.row == row), key=lambda item: item.column)


def _first_value_right(
    sheet: SheetGrid,
    row: int,
    column: int,
    *,
    excluded_labels: tuple[str, ...] = (),
    stop_labels: tuple[str, ...] = (),
) -> Any:
    excluded = tuple(_normalized_label(label) for label in excluded_labels)
    stops = tuple(_normalized_label(label) for label in stop_labels)
    for cell in _row_cells(sheet, row):
        if cell.column <= column:
            continue
        text = _clean_text(cell.value)
        if not text:
            continue
        normalized = _normalized_label(text)
        if stops and any(label in normalized for label in stops):
            return None
        if excluded and any(label in normalized for label in excluded):
            continue
        return cell.value
    return None


def _find_label_cell(sheet: SheetGrid, aliases: tuple[str, ...]) -> GridCell | None:
    normalized_aliases = tuple(_normalized_label(alias) for alias in aliases)
    for cell in sheet.non_empty():
        normalized = _normalized_label(cell.value)
        if any(alias in normalized for alias in normalized_aliases):
            return cell
    return None


def _value_after_label(
    sheet: SheetGrid,
    aliases: tuple[str, ...],
    *,
    allow_embedded: bool = True,
    stop_labels: tuple[str, ...] = (),
) -> Any:
    cell = _find_label_cell(sheet, aliases)
    if not cell:
        return None
    text = _clean_text(cell.value) or ""
    if allow_embedded:
        for separator in ("：", ":"):
            if separator in text:
                remainder = text.split(separator, 1)[1].strip()
                if remainder:
                    return remainder
    return _first_value_right(
        sheet,
        cell.row,
        cell.column,
        excluded_labels=aliases,
        stop_labels=stop_labels,
    )


def _parse_route(sheet: SheetGrid) -> tuple[str | None, str | None, str | None]:
    route_cell = _find_label_cell(sheet, ("航程", "voyage from"))
    if not route_cell:
        return None, None, None

    origin = None
    text = _clean_text(route_cell.value) or ""
    origin_match = re.search(r"自\s*(.+)$", text)
    if origin_match:
        origin = origin_match.group(1).strip()
    elif ":" in text or "：" in text:
        origin = re.split(r"[:：]", text, maxsplit=1)[1].strip() or None

    via = None
    destination = None
    row = _row_cells(sheet, route_cell.row)
    for index, cell in enumerate(row):
        normalized = _normalized_label(cell.value)
        if normalized in {_normalized_label("经"), _normalized_label("via")}:
            if index + 1 < len(row):
                candidate = _clean_text(row[index + 1].value)
                if candidate and _normalized_label(candidate) not in {
                    _normalized_label("至"),
                    _normalized_label("to"),
                }:
                    via = candidate
        if normalized in {_normalized_label("至"), _normalized_label("to")}:
            if index + 1 < len(row):
                destination = _clean_text(row[index + 1].value)

    if not destination:
        for cell in reversed(row):
            if cell.column > route_cell.column:
                candidate = _clean_text(cell.value)
                if candidate and _normalized_label(candidate) not in {
                    _normalized_label("经"),
                    _normalized_label("via"),
                    _normalized_label("至"),
                    _normalized_label("to"),
                }:
                    destination = candidate
                    break
    return origin, destination, via


def _parse_quantity(cell: GridCell | None) -> tuple[Decimal | None, str | None, Decimal | None]:
    if not cell:
        return None, None, None
    quantity = _decimal(cell.value)
    source = " ".join(filter(None, (_clean_text(cell.value), cell.number_format)))
    unit_match = re.search(r"\b(MT|KG|TONS?|TONNES?|CBM|PCS|BAGS?)\b", source, re.IGNORECASE)
    tolerance_match = re.search(r"(?:\+/-|±)\s*(\d+(?:\.\d+)?)\s*%", source)
    unit = unit_match.group(1).upper() if unit_match else None
    tolerance = Decimal(tolerance_match.group(1)) if tolerance_match else None
    return quantity, unit, tolerance


def _parse_cargo_rows(sheet: SheetGrid) -> list[dict[str, Any]]:
    header = _find_label_cell(sheet, ("保险货物项目", "items"))
    if not header:
        return []
    boundary = _find_label_cell(sheet, ("请将投保的险别", "risks insured against"))
    end_row = boundary.row if boundary else min(sheet.max_row + 1, header.row + 20)
    cargo_rows: list[dict[str, Any]] = []
    for row_number in range(header.row + 1, end_row):
        values = {cell.column: cell for cell in _row_cells(sheet, row_number)}
        normalized_row_values = {
            _normalized_label(cell.value)
            for cell in values.values()
            if _clean_text(cell.value)
        }
        header_values = {
            _normalized_label("items"),
            _normalized_label("保险货物项目"),
            _normalized_label("marks & nos"),
            _normalized_label("标记"),
            _normalized_label("quantities & packings"),
            _normalized_label("包装及数量"),
        }
        if normalized_row_values & header_values:
            continue
        cargo_cell = next((values[col] for col in range(1, 4) if col in values and _clean_text(values[col].value)), None)
        marks_cell = next((values[col] for col in range(4, 7) if col in values and _clean_text(values[col].value)), None)
        quantity_cell = next((values[col] for col in range(7, 10) if col in values and _clean_text(values[col].value)), None)
        cargo_text = _clean_text(cargo_cell.value) if cargo_cell else None
        if cargo_text and _normalized_label(cargo_text) in {
            _normalized_label("items"),
            _normalized_label("保险货物项目"),
        }:
            cargo_text = None
        quantity, unit, tolerance = _parse_quantity(quantity_cell)
        marks = _clean_text(marks_cell.value) if marks_cell else None
        if not any((cargo_text, marks, quantity is not None)):
            continue
        cargo_rows.append(
            {
                "cargo_name": cargo_text,
                "cargo_marks": marks,
                "quantity": quantity,
                "quantity_unit": unit,
                "quantity_tolerance_percent": tolerance,
            }
        )
    return cargo_rows


def _parse_vessel_information_sheet(sheet: SheetGrid) -> tuple[dict[str, Any], list[str]] | None:
    """Read the bilingual, row-oriented chartered-vessel information form."""
    aliases = {
        "business_type": ("业务类型",),
        "vessel_name": ("船名", "vessel name"),
        "imo_number": ("imo号", "imo number"),
        "flag": ("船旗", "flag"),
        "vessel_owner": ("船东", "owner of vessel"),
        "charterer": ("租船人", "charterer"),
        "classification": ("船级", "classification"),
        "pi_insurance": ("保赔险", "protection & indemnity insurance"),
        "built_year": ("船舶建造年份", "vessel built year"),
    }
    headers: dict[str, int] = {}
    header_row: int | None = None
    for row_number in range(1, min(sheet.max_row, 12) + 1):
        matches: dict[str, int] = {}
        for cell in _row_cells(sheet, row_number):
            label = _normalized_label(cell.value)
            for field, names in aliases.items():
                if label in {_normalized_label(name) for name in names}:
                    matches[field] = cell.column
        if {"vessel_name", "imo_number", "flag"} <= matches.keys():
            headers, header_row = matches, row_number
            break
    if header_row is None:
        return None

    shipments: list[dict[str, Any]] = []
    warnings: list[str] = []
    for row_number in range(header_row + 1, sheet.max_row + 1):
        vessel_name = _clean_text(sheet.value(row_number, headers["vessel_name"]))
        imo = _clean_text(sheet.value(row_number, headers["imo_number"]))
        if not vessel_name or _normalized_label(vessel_name) in {"vesselname", "船名"}:
            continue
        built_year = _decimal(sheet.value(row_number, headers.get("built_year", 0)))
        extra_data = {
            "sheet_name": sheet.name,
            "source_row": row_number,
            "form_type": "vessel_information",
        }
        for field in ("business_type", "vessel_owner", "charterer", "classification", "pi_insurance"):
            value = _clean_text(sheet.value(row_number, headers[field])) if field in headers else None
            if value:
                extra_data[field] = value
        shipments.append({
            "vessel_name": vessel_name,
            "imo_number": imo,
            "flag": _clean_text(sheet.value(row_number, headers["flag"])),
            "built_year": int(built_year) if built_year and 1800 <= built_year <= 2200 else None,
            "extra_data": extra_data,
        })
        if not imo:
            warnings.append(f"{sheet.name} 第 {row_number} 行：缺少 IMO 号")
    if not shipments:
        raise InsuranceWorkbookParseError("No vessel rows could be extracted from the workbook")
    warnings.append(f"{sheet.name}：船舶信息表未包含货物、保额或航次日期；请核实后以待审核状态保存")
    return {"insured_name": None, "received_at": None, "shipments": shipments, "extra_data": {"sheet_name": sheet.name}}, warnings


def _meaningful_text(value: Any) -> str | None:
    text = _clean_text(value)
    return None if text and text.casefold() in {"nil", "n/a", "na", "none"} else text


def _parse_underwriting_sheet(sheet: SheetGrid) -> tuple[dict[str, Any], list[str]] | None:
    headers = [_normalized_label(sheet.value(1, column)) for column in range(1, sheet.max_column + 1)]
    if not any(header.startswith("vesselnameaircraftnumber") for header in headers) or not any(
        header.startswith("vesselimonumber") for header in headers
    ):
        return None

    shipments: list[dict[str, Any]] = []
    for row_number in range(2, sheet.max_row + 1):
        vessel_name = _meaningful_text(sheet.value(row_number, 10))
        imo = _meaningful_text(sheet.value(row_number, 11))
        if not vessel_name or not imo:
            continue
        insured_amount = _meaningful_text(sheet.value(row_number, 14))
        invoice_amount = _meaningful_text(sheet.value(row_number, 13))
        currency_match = re.search(r"(USD|EUR|CNY|GBP|HKD)", insured_amount or invoice_amount or "", re.I)
        tolerance_match = re.search(r"(?:±|\+/-)\s*(\d+(?:\.\d+)?)\s*%", insured_amount or invoice_amount or "")
        extra_data: dict[str, Any] = {
            "sheet_name": sheet.name,
            "source_row": row_number,
            "form_type": "war_risk_underwriting",
        }
        row_insured = _meaningful_text(sheet.value(row_number, 2))
        if row_insured:
            extra_data["insured_name"] = row_insured
        extra_columns = {
            "declaration_sequence": 1,
            "cedant": 3,
            "etd_window": 4,
            "eta_window": 5,
            "intermediary_port_answer": 6,
            "transshipment_answer": 7,
            "other_route_description": 12,
            "hs_code": 16,
            "shipper": 17,
            "supplier": 18,
            "seller": 19,
            "member_assured": 20,
            "buyer": 21,
            "consignee": 22,
            "end_user": 23,
            "sanctions_screened": 24,
            "sanctions_findings": 25,
            "sanctions_licence": 26,
        }
        for field, column in extra_columns.items():
            value = _clean_text(sheet.value(row_number, column))
            if value:
                extra_data[field] = value
        if tolerance_match:
            extra_data["insured_value_tolerance_percent"] = tolerance_match.group(1)
        shipments.append({
            "vessel_name": vessel_name,
            "imo_number": imo,
            "etd": _datetime(sheet.value(row_number, 4)),
            "eta": _datetime(sheet.value(row_number, 5)),
            "loading_port": _meaningful_text(sheet.value(row_number, 8)),
            "discharge_port": _meaningful_text(sheet.value(row_number, 9)),
            "invoice_value": _decimal(invoice_amount),
            "insured_value": _decimal(insured_amount),
            "currency": currency_match.group(1).upper() if currency_match else None,
            "cargo_name": _meaningful_text(sheet.value(row_number, 15)),
            "extra_data": extra_data,
        })
    if not shipments:
        raise InsuranceWorkbookParseError("No voyage declarations could be extracted from the underwriting sheet")
    warnings = [f"{sheet.name}：航期区间按原文保存，请在核实环节确认具体 ETD/ETA"] if any(
        row["extra_data"].get("etd_window") and not row["etd"] for row in shipments
    ) else []
    return {
        "insured_name": _meaningful_text(sheet.value(2, 2)),
        "received_at": None,
        "shipments": shipments,
        "extra_data": {"sheet_name": sheet.name},
    }, warnings


def _parse_sheet(sheet: SheetGrid) -> tuple[dict[str, Any], list[str]]:
    warnings: list[str] = []
    insured_name = _clean_text(_value_after_label(sheet, ("被保险人", "insured")))
    vessel_name = _clean_text(_value_after_label(sheet, ("装载运输工具", "per conveyance"), allow_embedded=False))
    sailing = _datetime(_value_after_label(sheet, ("开航日期", "sailing on/about"), allow_embedded=False))
    application_date = _datetime(_value_after_label(sheet, ("投保日期", "applicant date"), allow_embedded=False))
    origin, destination, via = _parse_route(sheet)
    built_year = _decimal(_value_after_label(sheet, ("船龄",), allow_embedded=False))
    cargo_rows = _parse_cargo_rows(sheet) or [{}]

    common = {
        "vessel_name": vessel_name,
        "imo_number": _clean_text(_value_after_label(sheet, ("imo",), allow_embedded=False)),
        "voyage_number": _clean_text(_value_after_label(sheet, ("航次", "voyage no"), allow_embedded=False)),
        "built_year": int(built_year) if built_year and 1800 <= built_year <= 2200 else None,
        "invoice_number": _clean_text(
            _value_after_label(
                sheet,
                ("发票号", "invoice no"),
                allow_embedded=False,
                stop_labels=("提单号", "b/l no"),
            )
        ),
        "bill_of_lading_number": _clean_text(_value_after_label(sheet, ("提单号", "b/l no"), allow_embedded=False)),
        "invoice_value": _decimal(_value_after_label(sheet, ("发票金额", "invoice value"), allow_embedded=False)),
        "insured_value": _decimal(_value_after_label(sheet, ("保险金额", "amount insured"), allow_embedded=False)),
        "loading_port": origin,
        "discharge_port": destination,
        "via_port": via,
        "etd": sailing,
        "insurance_rate": _decimal(_value_after_label(sheet, ("费率", "rate"), allow_embedded=False)),
        "premium": _decimal(_value_after_label(sheet, ("保险费", "premium"), allow_embedded=False)),
    }
    shipments = [{**common, **cargo} for cargo in cargo_rows]
    if not vessel_name:
        warnings.append(f"{sheet.name}: vessel name was not detected")
    if not any(item.get("cargo_name") for item in shipments):
        warnings.append(f"{sheet.name}: cargo name was not detected")
    if not origin or not destination:
        warnings.append(f"{sheet.name}: route is incomplete")

    return {
        "insured_name": insured_name,
        "received_at": application_date,
        "shipments": shipments,
        "extra_data": {"sheet_name": sheet.name},
    }, warnings


def parse_declaration_workbook(file_name: str, content: bytes) -> dict[str, Any]:
    extension = Path(file_name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise InsuranceWorkbookParseError("Only .xls and .xlsx declaration files are supported")
    if not content:
        raise InsuranceWorkbookParseError("The declaration file is empty")

    sheets = _load_xls(content) if extension == ".xls" else _load_xlsx(content)
    if not sheets:
        raise InsuranceWorkbookParseError("The workbook has no worksheets")

    parsed_sheets: list[dict[str, Any]] = []
    warnings: list[str] = []
    for sheet in sheets:
        vessel_information = _parse_vessel_information_sheet(sheet)
        underwriting = _parse_underwriting_sheet(sheet) if not vessel_information else None
        if not vessel_information and not underwriting and "开口保单信息" in sheet.name:
            continue
        parsed, sheet_warnings = vessel_information or underwriting or _parse_sheet(sheet)
        if parsed["shipments"] and any(
            any(value is not None for key, value in shipment.items() if key != "extra_data")
            for shipment in parsed["shipments"]
        ):
            parsed_sheets.append(parsed)
        warnings.extend(sheet_warnings)

    if not parsed_sheets:
        raise InsuranceWorkbookParseError("No declaration data could be extracted from the workbook")

    insured_name = next((sheet["insured_name"] for sheet in parsed_sheets if sheet["insured_name"]), None)
    received_at = next((sheet["received_at"] for sheet in parsed_sheets if sheet["received_at"]), None)
    shipments: list[dict[str, Any]] = []
    for sheet in parsed_sheets:
        for shipment in sheet["shipments"]:
            shipment["extra_data"] = {
                **shipment.get("extra_data", {}),
                "sheet_name": sheet["extra_data"]["sheet_name"],
            }
            shipments.append(shipment)

    return {
        "parser_name": PARSER_NAME,
        "parser_version": PARSER_VERSION,
        "parsed_data": {
            "insured_name": insured_name,
            "received_at": received_at,
            "shipments": shipments,
            "extra_data": {"sheet_names": [sheet.name for sheet in sheets]},
        },
        "warnings": sorted(set(warnings)),
    }


def parse_declaration_materials(
    underwriting_file_name: str,
    underwriting_content: bytes,
    vessel_file_name: str,
    vessel_content: bytes,
) -> dict[str, Any]:
    underwriting = parse_declaration_workbook(underwriting_file_name, underwriting_content)
    vessel_information = parse_declaration_workbook(vessel_file_name, vessel_content)
    underwriting_rows = underwriting["parsed_data"]["shipments"]
    vessel_rows = vessel_information["parsed_data"]["shipments"]
    if not underwriting_rows or not all(row.get("extra_data", {}).get("form_type") == "war_risk_underwriting" for row in underwriting_rows):
        raise InsuranceWorkbookParseError("承保信息文件未识别到航次申报信息")
    if not vessel_rows or not all(row.get("extra_data", {}).get("form_type") == "vessel_information" for row in vessel_rows):
        raise InsuranceWorkbookParseError("船舶信息文件未识别到船舶资料")

    by_imo = {row["imo_number"]: row for row in vessel_rows if row.get("imo_number")}
    if len(by_imo) != len(vessel_rows):
        raise InsuranceWorkbookParseError("船舶信息表存在缺失或重复的 IMO，无法安全合并")
    merged_shipments = []
    matched_imos = set()
    for row in underwriting_rows:
        vessel = by_imo.get(row.get("imo_number"))
        if not vessel:
            raise InsuranceWorkbookParseError(f"承保信息中的 IMO {row.get('imo_number') or '缺失'} 未在船舶信息表中找到")
        matched_imos.add(row["imo_number"])
        merged_shipments.append({
            **vessel,
            **row,
            "extra_data": {
                **vessel.get("extra_data", {}),
                **row.get("extra_data", {}),
                "form_type": "two_form_declaration",
                "vessel_information_sheet": vessel.get("extra_data", {}).get("sheet_name"),
                "underwriting_sheet": row.get("extra_data", {}).get("sheet_name"),
            },
        })
    if matched_imos != set(by_imo):
        raise InsuranceWorkbookParseError("船舶信息表含有未出现在承保信息表中的 IMO，请核对两份材料")
    return {
        "parser_name": "marine_cargo_two_form",
        "parser_version": PARSER_VERSION,
        "parsed_data": {
            "insured_name": underwriting["parsed_data"].get("insured_name"),
            "received_at": underwriting["parsed_data"].get("received_at"),
            "shipments": merged_shipments,
            "extra_data": {
                "form_type": "two_form_declaration",
                "underwriting_file_name": underwriting_file_name,
                "vessel_file_name": vessel_file_name,
            },
        },
        "warnings": underwriting["warnings"],
    }
