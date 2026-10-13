"""Template-based OFAC report with declaration facts and embedded evidence."""
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from PIL import Image, ImageOps

REPORT_TEMPLATE_VERSION = 'sinochem-questionnaire-v1'
ASSETS = Path(__file__).parent / 'assets' / 'sanctions'
RESULTS = {'no_match': '未发现匹配', 'possible_match': '疑似匹配',
           'confirmed_match': '确认匹配', 'needs_escalation': '需升级复核'}


def _value(value):
    return '未提供' if value is None or str(value).strip() == '' else str(value)


def _answers(shipment, parties):
    extra = shipment.get('extra_data') or {}
    def field(name):
        return _value(shipment.get(name))
    def detail(name):
        return _value(extra.get(name))
    return [field('vessel_name'), field('imo_number'), field('flag'), detail('insured_name'),
            detail('vessel_owner'), detail('charterer'), detail('charter_party_period'),
            field('origin_country'), f"{field('cargo_name')}\nHS code: {detail('hs_code')}",
            detail('end_use'), field('loading_port'), field('discharge_port'),
            detail('loading_port_agent'), detail('discharge_port_agent'),
            f"货量：{field('quantity')} {shipment.get('quantity_unit') or ''}\n"
            f"发票金额：{field('invoice_value')} {shipment.get('currency') or ''}\n"
            f"保险金额：{field('insured_value')} {shipment.get('currency') or ''}",
            detail('shipper'), detail('supplier'), detail('buyer'), detail('seller'),
            detail('consignee'), detail('end_user'),
            '本报告仅记录 OFAC SDN 本地初筛及 OFAC 官网人工复核。未执行英国、联合国及欧盟名单筛查。',
            '\n'.join(f"{p['label']} {p['name']}：{RESULTS.get(p['review']['result'], p['review']['result'])}"
                      for p in parties) or '无已记录的主体复核结果', detail('sanctions_licence')]


def render_report(case, declaration, images, generated_at):
    document = Document(ASSETS / 'questionnaire-template.docx')
    timestamp = datetime.fromisoformat(generated_at.replace('Z', '+00:00')).astimezone(ZoneInfo('Asia/Shanghai'))
    header = document.sections[0].header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run().add_picture(str(ASSETS / 'sinochem-logo.png'), width=Inches(1.4))
    header.add_run(f"\nDate: {timestamp:%Y-%m-%d}").font.size = Pt(9)
    document.add_paragraph(f"申报编号：{case['declaration_number']}    企业：{case.get('company_name') or '未提供'}")
    document.add_paragraph(f"生成时间：{timestamp:%Y-%m-%d %H:%M:%S}（Asia/Shanghai）")
    document.add_paragraph('本报告记录系统初筛、OFAC 官网人工复核结论及对应截图。'
                           '本地初筛未发现候选不等于合规放行；本报告不自动认定制裁命中或合规放行。')
    base_table = document.tables[1]
    for index, shipment in enumerate(declaration.get('shipments') or []):
        if index:
            document.add_page_break()
            document.add_paragraph(f"风险信息 {index + 1}", style='Heading 1')
            table_xml = deepcopy(base_table._tbl)
            document._body._body.insert(-1, table_xml)
            table = document.tables[-1]
        else:
            table = base_table
        parties = [p for p in case['parties'] if p['shipment_id'] == shipment['id']]
        for row, value in zip(table.rows, _answers(shipment, parties)):
            row.cells[1].text = value
    for party, image in images:
        document.add_page_break()
        document.add_paragraph(f"{party['label']}  {party['name']}", style='Heading 1')
        document.add_paragraph(f"标的记录：{party['shipment_id']}")
        preliminary = party['preliminary']
        matches = preliminary['results']
        document.add_paragraph(f"系统初筛：{preliminary['screened_at']}；候选匹配 {len(matches)} 条")
        for match in matches:
            document.add_paragraph(f"{match.get('matched_name') or '—'}  相似度 {match.get('score') or 0}%  "
                                   f"OFAC ID {match.get('profile_id') or '—'}")
        if not matches:
            document.add_paragraph('未发现本地候选匹配，不等于合规放行。')
        review = party['review']
        document.add_paragraph(f"官网人工复核：{RESULTS.get(review['result'], review['result'])}；"
                               f"复核人：{review['reviewer']}；时间：{review['reviewed_at']}")
        document.add_paragraph(f"复核说明：{review.get('note') or '—'}")
        document.add_paragraph('OFAC 官方复核入口：https://sanctionssearch.ofac.treas.gov/')
        document.add_paragraph(f"截图原文件：{party['evidence'].get('file_name') or '—'}\n"
                               f"截图 SHA-256：{party['evidence']['sha256']}")
        with Image.open(BytesIO(image)) as source:
            screenshot = ImageOps.exif_transpose(source).convert('RGBA')
            background = Image.new('RGBA', screenshot.size, 'white')
            screenshot = Image.alpha_composite(background, screenshot).convert('RGB')
            normalized = BytesIO()
            screenshot.save(normalized, format='PNG')
            width, height = screenshot.size
        scale = min(6.0 / width, 6.5 / height)
        document.add_paragraph().add_run().add_picture(BytesIO(normalized.getvalue()),
            width=Inches(width * scale), height=Inches(height * scale))
    output = BytesIO()
    document.save(output)
    return output.getvalue()
