from datetime import datetime
from io import BytesIO
from zipfile import ZipFile

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from api.routes.insurance import (
    InsuranceRepositoryDependency,
    get_declaration_parser,
    get_declaration_materials_parser,
    get_insurance_repository,
    get_insurance_storage,
    get_insurance_tracking_service,
    get_insurance_submission_preparer,
    get_insurance_sanctions_workflow,
    router,
)
from risk_monitor.insurance_tracking_service import MmsiSourceUnavailableError, MmsiVerificationError
from risk_monitor.deepseek_mmsi_resolver import DeepSeekMmsiError


POLICY = {
    "id": "policy-1",
    "company_id": "company-1",
    "company_name": "中化化肥",
    "policy_name": "2026年度进出口货物运输保险",
    "policy_number": "PICC-2026-001",
    "insurer": "PICC",
    "start_date": "2026-01-01",
    "end_date": "2026-12-31",
    "currency": "USD",
    "status": "active",
}


class FakeRepository:
    def __init__(self):
        self.imports = {}
        self.declaration_status = "confirmed"

    def list_policies(self, **kwargs):
        return [POLICY]

    def get_policy(self, policy_id):
        return POLICY if policy_id == "policy-1" else None

    def create_declaration_import(self, payload):
        row = {
            **payload,
            "created_at": datetime(2026, 10, 8),
            "updated_at": datetime(2026, 10, 8),
        }
        self.imports[payload["id"]] = row
        return row

    def get_declaration_import(self, import_id):
        return self.imports.get(import_id)

    def list_pending_declaration_imports(self):
        return [row for row in self.imports.values() if row.get("status") == "pending_confirmation"]

    def confirm_declaration_import(self, **kwargs):
        self.confirmed_kwargs = kwargs
        return {"declaration_id": "declaration-1", "declaration_number": "DEC-2026-0001"}

    def confirm_declaration_import_batch(self, **kwargs):
        self.confirmed_kwargs = kwargs
        return [
            {"declaration_id": f"declaration-{index + 1}", "declaration_number": f"DEC-2026-{index + 1:04d}"}
            for index, _ in enumerate(kwargs["shipments"])
        ]

    def confirm_underwriting_terms(self, declaration_id, rate):
        self.terms = (declaration_id, rate)
        return {"shipment_id": "shipment-1", "premium": "720.52"}

    def confirm_underwriting(self, declaration_id):
        self.underwriting = declaration_id
        return {"shipment_id": "shipment-1", "premium": "720.52"}

    def submit_draft(self, declaration_id, screening):
        self.submitted_screening = screening
        self.declaration_status = "pending_review"

    def get_declaration(self, declaration_id):
        return {
            "id": declaration_id, "declaration_number": "DEC-2026-0001", "policy_id": "policy-1",
            "company_id": "company-1", "source_type": "manual_upload",
            "received_at": datetime(2026, 10, 11), "status": self.declaration_status,
            "shipments": [{"id": "shipment-1", "declaration_id": declaration_id,
                           "vessel_name": "Wadi Alkarnak", "imo_number": "9127136",
                           "insured_value": "720522", "currency": "USD", "insurance_rate": "0.1",
                           "premium": "720.52", "cargo_name": "Fertilizer", "loading_port": "Aqaba",
                           "discharge_port": "Huangpu", "extra_data": {"insured_name": "Sinochem"}}],
        }

    def list_insured_vessels(self):
        return [{"vessel_name": "Wadi Alkarnak", "imo_number": "9127136", "declaration_count": 1, "declaration_number": "DEC-2026-0001", "company_name": "中化化肥"}]

    def update_declaration_status(self, declaration_id, status):
        if declaration_id != "declaration-1":
            return None
        return {
            "id": declaration_id,
            "declaration_number": "DEC-2026-0001",
            "policy_id": "policy-1",
            "company_id": "company-1",
            "source_type": "manual_upload",
            "received_at": datetime(2026, 10, 8),
            "status": status,
        }

    def list_documents(self, **kwargs):
        return []


class FakeStorage:
    bucket = "insurance-documents"

    def __init__(self):
        self.uploaded = {}

    def build_path(self, owner_type, owner_id, file_name):
        return f"{owner_type}/{owner_id}/{file_name}"

    def upload(self, *, path, content, content_type):
        self.uploaded[path] = content

    def remove(self, path):
        self.uploaded.pop(path, None)

    def signed_url(self, path, expires_in=3600):
        return f"https://example.test/{path}"


class FakeTrackingService:
    def __init__(self):
        self.shipments = []
        self.candidate = {"mmsi": "622121425", "source_url": "https://example.test/vessel/9127136", "source_title": "Vessel identity"}
        self.verification = "match"

    def suggest_mmsi(self, imo, vessel_name):
        return self.candidate

    def verify_confirmed_mmsi(self, imo, mmsi):
        if self.verification == "mismatch":
            raise MmsiVerificationError("MMSI does not match the declared IMO in Chinaports.")
        if self.verification == "unavailable":
            raise MmsiSourceUnavailableError("Chinaports identity lookup is unavailable.")
        return {"imo": imo, "mmsi": mmsi}

    def validate_confirmed_shipments(self, shipments):
        if not any((item.get("extra_data") or {}).get("mmsi_confirmed") for item in shipments):
            return
        if self.verification == "unavailable":
            raise MmsiSourceUnavailableError("Chinaports identity lookup is unavailable.")
        if self.verification == "mismatch":
            raise MmsiVerificationError("MMSI does not match the declared IMO in Chinaports.")

    def sync_shipments(self, shipments):
        self.shipments.extend(shipments)

    def status_for(self, imo, confirmed_mmsi=None):
        return "awaiting_ais_provider"


class FakeSanctionsWorkflow:
    def __init__(self):
        self.ready = True
        self.terms_ready = False
        self.checked = []

    def assert_underwriting_ready(self, declaration_id):
        self.checked.append(declaration_id)
        if not self.ready:
            raise ValueError("请先完成全部主体筛查、上传截图并生成筛查报告。")
        return {"underwriting": {"status": "confirmed" if self.terms_ready else "pending"}}

    def prepare_atomic_confirmation(self, declaration_id):
        self.assert_underwriting_ready(declaration_id)
        return {'workflow': {}, 'sha256': 'verified'}


def _workbook_bytes():
    workbook = Workbook()
    workbook.active["A1"] = "test"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _parser(file_name, content):
    return {
        "parser_name": "fake",
        "parser_version": "1",
        "parsed_data": {
            "insured_name": "中化化肥",
            "received_at": "2026-10-08T00:00:00",
            "shipments": [
                {
                    "vessel_name": "PACIFIC PROVIDENCE",
                    "cargo_name": "MOP",
                    "quantity": "50000",
                    "quantity_unit": "MT",
                }
            ],
        },
        "warnings": [],
    }


def _client():
    app = FastAPI()
    app.include_router(router)
    repository = FakeRepository()
    storage = FakeStorage()
    tracking_service = FakeTrackingService()
    sanctions_workflow = FakeSanctionsWorkflow()
    app.dependency_overrides[get_insurance_repository] = lambda: repository
    app.dependency_overrides[get_insurance_storage] = lambda: storage
    app.dependency_overrides[get_declaration_parser] = lambda: _parser
    app.dependency_overrides[get_declaration_materials_parser] = lambda: (lambda *args: _parser(args[0], args[1]))
    app.dependency_overrides[get_insurance_tracking_service] = lambda: tracking_service
    from risk_monitor.insurance_screening import prepare_shipments_for_screening
    app.dependency_overrides[get_insurance_submission_preparer] = lambda: prepare_shipments_for_screening
    app.dependency_overrides[get_insurance_sanctions_workflow] = lambda: sanctions_workflow
    repository.tracking_service = tracking_service
    repository.sanctions_workflow = sanctions_workflow
    return TestClient(app), repository, storage


def test_two_form_upload_preserves_both_originals_in_one_import():
    client, repository, storage = _client()
    response = client.post(
        "/api/v1/insurance/declaration-imports/bundle",
        data={"policy_id": "policy-1"},
        files={
            "underwriting_file": ("underwriting.xlsx", _workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            "vessel_file": ("vessel.xlsx", _workbook_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        },
    )
    assert response.status_code == 201
    assert response.json()["source_file_name"] == "投保申报材料.zip"
    assert len(repository.imports) == 1
    with ZipFile(BytesIO(next(iter(storage.uploaded.values())))) as bundle:
        assert sorted(bundle.namelist()) == ["承保信息/underwriting.xlsx", "船舶信息/vessel.xlsx"]
        assert load_workbook(BytesIO(bundle.read("承保信息/underwriting.xlsx"))).active["A1"].value == "test"


def test_pending_import_can_be_listed_for_resume():
    client, _, _ = _client()
    created = client.post(
        "/api/v1/insurance/declaration-imports/bundle",
        data={"policy_id": "policy-1"},
        files={
            "underwriting_file": ("underwriting.xlsx", _workbook_bytes()),
            "vessel_file": ("vessel.xlsx", _workbook_bytes()),
        },
    )
    assert created.status_code == 201
    listed = client.get("/api/v1/insurance/declaration-imports")
    assert listed.status_code == 200
    assert [row["id"] for row in listed.json()] == [created.json()["id"]]


def test_lists_policy_summaries():
    client, _, _ = _client()

    response = client.get("/api/v1/insurance/policies")

    assert response.status_code == 200
    assert response.json()[0]["policy_number"] == "PICC-2026-001"


def test_manual_vessel_information_creates_pending_declaration():
    client, repository, storage = _client()
    response = client.post("/api/v1/insurance/manual-declarations", json={
        "policy_id": "policy-1", "insured_name": "Sinochem",
        "shipment": {"vessel_name": "Wadi Alkarnak", "imo_number": "9127136", "built_year": 1997,
                     "extra_data": {"form_type": "vessel_information", "charterer": "ARAB POTASH COMPANY"}},
    })
    assert response.status_code == 201
    assert repository.confirmed_kwargs["status"] == "pending_review"
    assert repository.imports[repository.confirmed_kwargs["import_id"]]["parsed_data"]["shipments"][0]["extra_data"]["charterer"] == "ARAB POTASH COMPANY"
    assert storage.uploaded


def test_manual_two_form_fields_are_kept_in_parsed_data():
    client, repository, _ = _client()
    response = client.post("/api/v1/insurance/manual-declarations", json={
        "policy_id": "policy-1", "insured_name": "Sinochem Fertilizer Macao Limited",
        "received_at": "2026-10-11T00:00:00+08:00",
        "shipment": {
            "vessel_name": "Wadi Alkarnak", "imo_number": "9127136",
            "cargo_name": "Muriate of Potash", "insured_value": "21356500",
            "currency": "USD", "loading_port": "Aqaba", "discharge_port": "China main port(s)",
            "extra_data": {"form_type": "two_form_declaration", "etd_window": "20-24 September,2026"},
        },
    })
    assert response.status_code == 201
    assert repository.confirmed_kwargs["status"] == "pending_review"
    parsed = repository.imports[repository.confirmed_kwargs["import_id"]]["parsed_data"]
    assert parsed["insured_name"] == "Sinochem Fertilizer Macao Limited"
    assert parsed["received_at"] == "2026-10-11T00:00:00+08:00"
    assert parsed["shipments"][0]["extra_data"]["etd_window"] == "20-24 September,2026"
    assert repository.tracking_service.shipments == []
    assert parsed["shipments"][0]["extra_data"]["sanctions_screening"]["status"] == "pending"


def test_insured_vessels_are_available_to_other_modules():
    client, _, _ = _client()
    response = client.get("/api/v1/insurance/insured-vessels")
    assert response.status_code == 200
    assert response.json()[0]["imo_number"] == "9127136"


def test_uploads_and_parses_declaration_before_confirmation():
    client, repository, storage = _client()

    response = client.post(
        "/api/v1/insurance/declaration-imports",
        data={"policy_id": "policy-1", "source_type": "manual_upload"},
        files={
            "file": (
                "2026 declaration.xlsx",
                _workbook_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )

    assert response.status_code == 201
    body = response.json()
    assert body["status"] == "pending_confirmation"
    assert body["parsed_data"]["shipments"][0]["vessel_name"] == "PACIFIC PROVIDENCE"
    assert repository.imports
    assert storage.uploaded


def test_confirms_import_and_returns_declaration_number():
    client, repository, _ = _client()
    repository.imports["import-1"] = {"parsed_data": {"insured_name": "Sinochem", "shipments": [{"imo_number": "9127136", "extra_data": {}}]}}

    response = client.post(
        "/api/v1/insurance/declaration-imports/import-1/confirm",
        json={"status": "pending_review"},
    )

    assert response.status_code == 201
    assert response.json()["declaration_number"] == "DEC-2026-0001"
    assert response.json()["declarations"] == []


def test_mmsi_suggestion_exposes_web_evidence_and_identity_check():
    client, _, _ = _client()
    response = client.post("/api/v1/insurance/mmsi-suggestion", json={"imo_number": "9127136", "vessel_name": "Wadi Alkarnak"})
    assert response.status_code == 200
    assert response.json()["mmsi"] == "622121425"
    assert response.json()["source"] == "deepseek_web_search"
    assert response.json()["source_url"] == "https://example.test/vessel/9127136"
    assert response.json()["identity_check"] == "matched"


def test_mmsi_suggestion_marks_wrong_candidate_before_confirmation():
    client, repository, _ = _client()
    repository.tracking_service.verification = "mismatch"
    response = client.post("/api/v1/insurance/mmsi-suggestion", json={"imo_number": "9127136", "vessel_name": "Wadi Alkarnak"})
    assert response.status_code == 200
    assert response.json()["identity_check"] == "mismatch"


def test_mmsi_suggestion_reports_missing_web_search_configuration():
    client, repository, _ = _client()
    def unavailable(imo, vessel_name):
        raise DeepSeekMmsiError("未配置 BRAVE_SEARCH_API_KEY，DeepSeek 无法联网检索 MMSI。")
    repository.tracking_service.suggest_mmsi = unavailable
    response = client.post("/api/v1/insurance/mmsi-suggestion", json={"imo_number": "9127136"})
    assert response.status_code == 503
    assert "BRAVE_SEARCH_API_KEY" in response.json()["detail"]


def test_mmsi_verification_endpoint_accepts_matching_manual_number():
    client, _, _ = _client()
    response = client.post("/api/v1/insurance/mmsi-verification", json={"imo_number": "9127136", "mmsi": "622121425"})
    assert response.status_code == 200
    assert response.json()["mmsi"] == "622121425"


def test_mmsi_verification_endpoint_explains_mismatch_and_source_failure():
    client, repository, _ = _client()
    repository.tracking_service.verification = "mismatch"
    mismatch = client.post("/api/v1/insurance/mmsi-verification", json={"imo_number": "9127136", "mmsi": "622123456"})
    assert mismatch.status_code == 422
    repository.tracking_service.verification = "unavailable"
    unavailable = client.post("/api/v1/insurance/mmsi-verification", json={"imo_number": "9127136", "mmsi": "622121425"})
    assert unavailable.status_code == 503


def test_manual_confirmed_mmsi_rejects_imo_mismatch_before_persistence():
    client, repository, storage = _client()
    repository.tracking_service.verification = "mismatch"
    response = client.post("/api/v1/insurance/manual-declarations", json={
        "policy_id": "policy-1", "insured_name": "Sinochem",
        "shipment": {"vessel_name": "Wadi Alkarnak", "imo_number": "9127136", "extra_data": {"mmsi_candidate": "622121425", "mmsi_confirmed": True}},
    })
    assert response.status_code == 422
    assert not repository.imports
    assert not storage.uploaded


def test_manual_confirmed_mmsi_does_not_start_tracking_before_underwriting():
    client, repository, _ = _client()
    response = client.post("/api/v1/insurance/manual-declarations", json={
        "policy_id": "policy-1", "insured_name": "Sinochem",
        "shipment": {"vessel_name": "Wadi Alkarnak", "imo_number": "9127136", "extra_data": {"mmsi_candidate": "622121425", "mmsi_confirmed": True}},
    })
    assert response.status_code == 201
    assert repository.tracking_service.shipments == []


def test_confirmed_import_rejects_unavailable_identity_source():
    client, repository, _ = _client()
    repository.tracking_service.verification = "unavailable"
    repository.imports["import-1"] = {"parsed_data": {"shipments": [{"imo_number": "9127136", "extra_data": {"mmsi_candidate": "622121425", "mmsi_confirmed": True}}]}}
    response = client.post("/api/v1/insurance/declaration-imports/import-1/confirm", json={"status": "pending_review"})
    assert response.status_code == 503
    assert not hasattr(repository, "confirmed_kwargs")


def test_multiline_import_creates_one_declaration_per_row_without_tracking():
    client, repository, _ = _client()
    repository.imports["import-1"] = {"parsed_data": {"insured_name": "Sinochem", "shipments": [
        {"vessel_name": "SAME SHIP", "imo_number": "9127136", "extra_data": {"source_row": 2, "declaration_sequence": "1"}},
        {"vessel_name": "SAME SHIP", "imo_number": "9127136", "extra_data": {"source_row": 3, "declaration_sequence": "2"}},
    ]}}
    response = client.post("/api/v1/insurance/declaration-imports/import-1/confirm", json={"status": "pending_review"})
    assert response.status_code == 201
    assert len(response.json()["declarations"]) == 2
    assert [item["extra_data"]["source_row"] for item in repository.confirmed_kwargs["shipments"]] == [2, 3]
    assert repository.tracking_service.shipments == []


def test_underwriting_confirmation_starts_tracking_after_commit():
    client, repository, _ = _client()
    response = client.post("/api/v1/insurance/declarations/declaration-1/underwriting-terms", json={"rate": "0.1"})
    assert response.status_code == 200
    assert repository.terms == ("declaration-1", "0.1")
    assert repository.tracking_service.shipments == []
    repository.sanctions_workflow.terms_ready = True
    response = client.post("/api/v1/insurance/declarations/declaration-1/underwriting-confirm")
    assert response.status_code == 200
    assert repository.underwriting == "declaration-1"
    assert repository.sanctions_workflow.checked == ["declaration-1", "declaration-1"]
    assert repository.tracking_service.shipments[0]["imo_number"] == "9127136"


def test_underwriting_is_blocked_until_report_is_complete():
    client, repository, _ = _client()
    repository.sanctions_workflow.ready = False
    response = client.post("/api/v1/insurance/declarations/declaration-1/underwriting-terms", json={"rate": "0.1"})
    assert response.status_code == 409
    assert "筛查报告" in response.json()["detail"]
    assert not hasattr(repository, "underwriting")
    assert repository.tracking_service.shipments == []


def test_final_confirmation_requires_terms_even_with_screening_report():
    client, repository, _ = _client()
    response = client.post("/api/v1/insurance/declarations/declaration-1/underwriting-confirm")
    assert response.status_code == 409
    assert "承保条件" in response.json()["detail"]
    assert not hasattr(repository, "underwriting")


def test_atomic_confirmation_is_generic_and_starts_tracking_once():
    client, repository, _ = _client()
    repository.declaration_status = 'pending_review'
    calls = []
    def confirm(declaration_id, rate, verification):
        calls.append((declaration_id, rate))
        applied = repository.declaration_status != 'confirmed'
        repository.declaration_status = 'confirmed'
        return {'applied': applied}
    repository.confirm_underwriting_atomic = confirm
    for _ in range(2):
        response = client.post('/api/v1/insurance/declarations/new-independent-id/underwriting-complete', json={'rate': '0.1'})
        assert response.status_code == 200
        assert response.json()['status'] == 'confirmed'
    assert calls == [('new-independent-id', '0.1')] * 2
    assert len(repository.tracking_service.shipments) == 1


def test_atomic_failure_never_starts_tracking():
    client, repository, _ = _client()
    repository.declaration_status = 'pending_review'
    def confirm(*args):
        raise ValueError('报告已过期')
    repository.confirm_underwriting_atomic = confirm
    response = client.post('/api/v1/insurance/declarations/another-id/underwriting-complete', json={'rate': 1})
    assert response.status_code == 409
    assert repository.declaration_status == 'pending_review'
    assert repository.tracking_service.shipments == []
    for rate in [-1, 101, 'NaN', 'Infinity', '']:
        assert client.post('/api/v1/insurance/declarations/another-id/underwriting-complete', json={'rate': rate}).status_code == 422
    assert repository.tracking_service.shipments == []


def test_draft_submit_queues_parties_without_starting_tracking():
    client, repository, _ = _client()
    repository.declaration_status = "draft"
    response = client.post("/api/v1/insurance/declarations/declaration-1/submit")
    assert response.status_code == 200
    assert response.json()["status"] == "pending_review"
    assert repository.submitted_screening["status"] == "pending"
    assert repository.tracking_service.shipments == []


def test_import_confirmation_persists_the_validated_shipment_snapshot():
    client, repository, _ = _client()
    repository.imports["import-1"] = {"parsed_data": {"insured_name": "Sinochem", "shipments": [{"imo_number": "9127136", "extra_data": {"mmsi_candidate": "622121425", "mmsi_confirmed": True}}]}}
    response = client.post("/api/v1/insurance/declaration-imports/import-1/confirm", json={"status": "pending_review"})
    assert response.status_code == 201
    assert repository.confirmed_kwargs["shipments"][0]["extra_data"]["mmsi_candidate"] == "622121425"
    assert repository.tracking_service.shipments == []


def test_rejects_unknown_policy_during_import():
    client, _, _ = _client()

    response = client.post(
        "/api/v1/insurance/declaration-imports",
        data={"policy_id": "missing"},
        files={"file": ("declaration.xlsx", _workbook_bytes())},
    )

    assert response.status_code == 404


def test_voids_only_the_selected_declaration():
    client, _, _ = _client()

    response = client.patch(
        "/api/v1/insurance/declarations/declaration-1",
        json={"status": "cancelled"},
    )

    assert response.status_code == 200
    assert response.json()["id"] == "declaration-1"
    assert response.json()["status"] == "cancelled"
    assert client.patch(
        "/api/v1/insurance/declarations/missing",
        json={"status": "cancelled"},
    ).status_code == 404
    assert client.patch(
        "/api/v1/insurance/declarations/declaration-1",
        json={"status": "confirmed"},
    ).status_code == 422


def test_repository_dependency_initializes_without_recursion(monkeypatch):
    class Repository:
        def get_dashboard_summary(self):
            return {"active_policy_count": 0}

    repository = Repository()
    monkeypatch.setattr(
        "risk_monitor.insurance_repository.InsuranceRepository",
        lambda: repository,
    )

    dependency = InsuranceRepositoryDependency()

    assert dependency.get_dashboard_summary() == {"active_policy_count": 0}
    assert dependency.__dict__["_repository"] is repository
