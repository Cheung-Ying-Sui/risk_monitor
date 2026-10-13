from io import BytesIO

from openpyxl import Workbook

from risk_monitor.insurance_excel_parser import parse_declaration_materials, parse_declaration_workbook


def _sinochem_workbook() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "约钾"
    sheet["A8"] = "被 保 险 人：SINOCHEM FERTILIZER MACAO LTD."
    sheet["A12"] = "发 票 号："
    sheet["F12"] = "提 单 号："
    sheet["A14"] = "装 载 运 输 工 具："
    sheet["C14"] = "PACIFIC PROVIDENCE"
    sheet["F14"] = "开航日期："
    sheet["G14"] = "2026-08-10"
    sheet["A17"] = "航 程：自 Aqaba,Jordan"
    sheet["E17"] = "经"
    sheet["F17"] = "至"
    sheet["G17"] = "China main port(s)"
    sheet["A20"] = "船龄："
    sheet["B20"] = 2025
    sheet["A26"] = "保险货物项目"
    sheet["E26"] = "标记"
    sheet["G26"] = "包装及数量"
    sheet["A27"] = "Items"
    sheet["D27"] = "Marks & Nos"
    sheet["G27"] = "Quantities&Packings"
    sheet["B29"] = "MOP"
    sheet["E29"] = "N/M"
    sheet["G29"] = 50000
    sheet["G29"].number_format = '#,##0" MT +/-10%"'
    sheet["A31"] = "请将投保的险别及条件注明如下："
    sheet["F39"] = "投保日期："
    sheet["H39"] = "2026-07-21"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_parses_known_marine_cargo_form_structure():
    result = parse_declaration_workbook("declaration.xlsx", _sinochem_workbook())

    assert result["parser_name"] == "marine_cargo_excel"
    parsed = result["parsed_data"]
    assert parsed["insured_name"] == "SINOCHEM FERTILIZER MACAO LTD."
    assert parsed["received_at"].date().isoformat() == "2026-07-21"
    assert len(parsed["shipments"]) == 1
    shipment = parsed["shipments"][0]
    assert shipment["vessel_name"] == "PACIFIC PROVIDENCE"
    assert shipment["invoice_number"] is None
    assert shipment["bill_of_lading_number"] is None
    assert shipment["cargo_name"] == "MOP"
    assert shipment["cargo_marks"] == "N/M"
    assert shipment["quantity"] == 50000
    assert shipment["quantity_unit"] == "MT"
    assert shipment["quantity_tolerance_percent"] == 10
    assert shipment["loading_port"] == "Aqaba,Jordan"
    assert shipment["discharge_port"] == "China main port(s)"
    assert shipment["etd"].date().isoformat() == "2026-08-10"
    assert shipment["built_year"] == 2025


def test_rejects_unsupported_file_type():
    try:
        parse_declaration_workbook("declaration.csv", b"x")
    except ValueError as exc:
        assert "Only .xls and .xlsx" in str(exc)
    else:
        raise AssertionError("Unsupported file type should fail")


def test_parses_chartered_vessel_information_form():
    workbook = Workbook()
    sheet = workbook.active
    sheet["A2"] = "业务类型"
    sheet["B2"] = "船名"
    sheet["C2"] = "IMO号"
    sheet["D2"] = "船旗"
    sheet["E2"] = "船东"
    sheet["F2"] = "租船人"
    sheet["G2"] = "船级"
    sheet["H2"] = "保赔险"
    sheet["I2"] = "船舶建造年份"
    sheet.append([])
    for column, value in enumerate(("约旦经红海", "Wadi Alkarnak", "9127136", "EGYPTIAN", "NATIONAL NAVIGATION COMPANY", "ARAB POTASH COMPANY", "LR", None, 1997), start=1):
        sheet.cell(4, column).value = value
    buffer = BytesIO()
    workbook.save(buffer)

    result = parse_declaration_workbook("vessel.xlsx", buffer.getvalue())
    shipment = result["parsed_data"]["shipments"][0]
    assert shipment["imo_number"] == "9127136"
    assert shipment["flag"] == "EGYPTIAN"
    assert shipment["built_year"] == 1997
    assert shipment["extra_data"]["business_type"] == "约旦经红海"
    assert shipment["extra_data"]["vessel_owner"] == "NATIONAL NAVIGATION COMPANY"
    assert shipment["extra_data"]["charterer"] == "ARAB POTASH COMPANY"
    assert shipment["extra_data"]["classification"] == "LR"
    assert shipment["extra_data"]["form_type"] == "vessel_information"
    assert any("待审核" in warning for warning in result["warnings"])


def _two_form_workbooks(imo="9127136"):
    underwriting = Workbook()
    voyage = underwriting.active
    voyage.title = "单航次或开口保单申报航次信息"
    voyage["J1"] = "Vessel Name/Aircraft Number\n船舶名称/航空器编号"
    voyage["K1"] = "Vessel IMO number\n船舶IMO号"
    voyage["B2"] = "Sinochem Fertilizer Macao Limited"
    voyage["D2"] = "20-24 September,2026"
    voyage["E2"] = "14-18 October,2026"
    voyage["H2"] = "Aqaba"
    voyage["I2"] = "China main port(s)"
    voyage["J2"] = "Wadi Alkarnak"
    voyage["K2"] = imo
    voyage["N2"] = "USD21,356,500.00±10%"
    voyage["O2"] = "Muriate of Potash"
    voyage["Q2"] = "THE ARAB POTASH CO., PLC."
    voyage["V2"] = "Sinochem Fertilizer Macao Limited"
    underwriting_bytes = BytesIO()
    underwriting.save(underwriting_bytes)

    vessels = Workbook()
    sheet = vessels.active
    for column, value in enumerate(("业务类型", "船名", "IMO号", "船旗", "船东", "租船人", "船级", "保赔险", "船舶建造年份"), start=1):
        sheet.cell(2, column).value = value
    for column, value in enumerate(("约旦经红海", "Wadi Alkarnak", "9127136", "EGYPTIAN", "NATIONAL NAVIGATION COMPANY", "ARAB POTASH COMPANY", "LR", None, 1997), start=1):
        sheet.cell(4, column).value = value
    vessel_bytes = BytesIO()
    vessels.save(vessel_bytes)
    return underwriting_bytes.getvalue(), vessel_bytes.getvalue()


def test_merges_two_declaration_materials_by_imo():
    underwriting, vessels = _two_form_workbooks()
    result = parse_declaration_materials("underwriting.xlsx", underwriting, "vessels.xlsx", vessels)
    assert result["parsed_data"]["insured_name"] == "Sinochem Fertilizer Macao Limited"
    shipment = result["parsed_data"]["shipments"][0]
    assert shipment["imo_number"] == "9127136"
    assert shipment["flag"] == "EGYPTIAN"
    assert shipment["cargo_name"] == "Muriate of Potash"
    assert str(shipment["insured_value"]) == "21356500.00"
    assert shipment["currency"] == "USD"
    assert shipment["etd"] is None
    assert shipment["extra_data"]["etd_window"] == "20-24 September,2026"
    assert shipment["extra_data"]["insured_value_tolerance_percent"] == "10"
    assert shipment["extra_data"]["form_type"] == "two_form_declaration"


def test_two_form_import_rejects_unmatched_imo():
    underwriting, vessels = _two_form_workbooks(imo="9999999")
    try:
        parse_declaration_materials("underwriting.xlsx", underwriting, "vessels.xlsx", vessels)
    except ValueError as exc:
        assert "9999999" in str(exc)
    else:
        raise AssertionError("Mismatched IMO should fail")


def test_underwriting_sheet_keeps_each_excel_row_even_for_same_imo():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "单航次或开口保单申报航次信息"
    sheet["J1"] = "Vessel Name/Aircraft Number\n船舶名称/航空器编号"
    sheet["K1"] = "Vessel IMO number\n船舶IMO号"
    for row, sequence, invoice in ((2, 1, "USD655,020.00±10%"), (3, 2, "USD332,910.00±10%")):
        sheet.cell(row, 1).value = sequence
        sheet.cell(row, 2).value = "Sinochem Fertilizer Macao Limited"
        sheet.cell(row, 10).value = "CMA CGM PALAIS ROYAL"
        sheet.cell(row, 11).value = "9839181"
        sheet.cell(row, 13).value = invoice
        sheet.cell(row, 14).value = invoice
        sheet.cell(row, 15).value = "Fertilizer"
    buffer = BytesIO()
    workbook.save(buffer)
    parsed = parse_declaration_workbook("underwriting.xlsx", buffer.getvalue())["parsed_data"]["shipments"]
    assert len(parsed) == 2
    assert [row["extra_data"]["declaration_sequence"] for row in parsed] == ["1", "2"]
    assert [str(row["invoice_value"]) for row in parsed] == ["655020.00", "332910.00"]
    assert all(row["extra_data"]["insured_name"] == "Sinochem Fertilizer Macao Limited" for row in parsed)
